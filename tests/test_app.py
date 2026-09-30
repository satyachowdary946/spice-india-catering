import os
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from pathlib import Path
from urllib.parse import unquote

TEST_DB = Path(__file__).parent / "test_catering.db"
if TEST_DB.exists():
    TEST_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB}"
os.environ["SESSION_SECRET"] = "test-secret-only-for-tests"

from fastapi.testclient import TestClient
from app.main import app, format_order_number, parse_order_sequence
from app.db import SessionLocal
from app.models import BusinessSettings


def csrf_from(html: str) -> str:
    marker = 'name="csrf_token" value="'
    return html.split(marker, 1)[1].split('"', 1)[0]


def quote_payload(phone="0891234567", event_name="Test Event", days=10):
    return {
        "details": {
            "name": "Test Customer",
            "phone": phone,
            "whatsapp": phone,
            "email": "customer@example.com",
            "event_date": str(datetime.now(ZoneInfo("Europe/Dublin")).date() + timedelta(days=days)),
            "event_day": "ignored-server-side",
            "event_name": event_name,
            "event_time": "18:30",
            "delivery_time": "17:45",
            "adults": "20",
            "kids": "4",
            "address": "Test address",
            "eircode": "N37K5P3",
        },
        "item_ids": [1, 2],
        "requested_dishes": ["Special Paneer Dish"],
        "customer_notes": "Please keep one section mild for children.",
    }


def test_build8_mobile_workflow():
    assert format_order_number(1) == "CAT0001"
    assert format_order_number(1000) == "CAT1000"
    assert format_order_number(1001) == "CAT10001"
    assert format_order_number(2001) == "CAT20001"
    assert parse_order_sequence("CAT0001") == 1
    assert parse_order_sequence("CAT1000") == 1000
    assert parse_order_sequence("CAT10001") == 1001
    assert parse_order_sequence("CAT20001") == 2001

    with TestClient(app) as client:
        assert client.get("/health").json()["status"] == "ok"
        home = client.get("/")
        assert home.status_code == 200
        assert '>Home<' in home.text

        menu_page = client.get("/menu")
        assert menu_page.status_code == 200
        assert 'data-exit-search' in menu_page.text
        assert 'data-request-rows' in menu_page.text
        assert 'data-add-request-row' in menu_page.text
        # The page is already the menu, so the primary header does not repeat a Menu link.
        header_html = menu_page.text.split('</header>', 1)[0]
        assert 'href="/menu">Menu</a>' not in header_html

        setup = client.get("/admin/setup")
        csrf = csrf_from(setup.text)
        r = client.post("/admin/setup", data={
            "csrf_token": csrf,
            "email": "owner@example.com",
            "password": "StrongPass123!",
            "confirm_password": "StrongPass123!",
        }, follow_redirects=False)
        assert r.status_code == 303

        # Strict phone and Eircode validation.
        bad = quote_payload(phone="08912abc67")
        response = client.post("/api/quotes", json=bad)
        assert response.status_code == 422
        assert "digits only" in response.text
        bad = quote_payload()
        bad["details"]["eircode"] = "N37BAD"
        response = client.post("/api/quotes", json=bad)
        assert response.status_code == 422
        assert "exactly 7" in response.text

        too_soon = quote_payload(days=2)
        response = client.post("/api/quotes", json=too_soon)
        assert response.status_code == 422
        assert "two full days" in response.text

        q1 = client.post("/api/quotes", json=quote_payload(event_name="Wedding"))
        assert q1.status_code == 200, q1.text
        b1 = q1.json()
        assert b1["order_number"] == "CAT0001"

        q2 = client.post("/api/quotes", json=quote_payload(event_name="Birthday", days=12))
        assert q2.status_code == 200, q2.text
        b2 = q2.json()
        assert b2["order_number"] == "CAT0002"

        # Regression: if the stored sequence falls behind, existing CAT numbers are skipped.
        with SessionLocal() as db:
            biz = db.get(BusinessSettings, 1)
            biz.next_order_sequence = 1
            db.commit()
        q3 = client.post("/api/quotes", json=quote_payload(event_name="Sequence Collision Test", days=14))
        assert q3.status_code == 200, q3.text
        assert q3.json()["order_number"] == "CAT0003"

        # Tracking works with either identifier. Phone with several orders gives a chooser.
        tracked = client.post("/track", data={"order_number": b1["order_number"], "phone": ""}, follow_redirects=False)
        assert tracked.status_code == 303
        assert tracked.headers["location"].endswith("/orders/" + b1["token"])
        tracked_phone = client.post("/track", data={"order_number": "", "phone": "0891234567"}, follow_redirects=False)
        assert tracked_phone.status_code == 200
        assert "Choose Your Order" in tracked_phone.text
        assert "Wedding" in tracked_phone.text and "Birthday" in tracked_phone.text
        assert "Delivery" in tracked_phone.text

        orders = client.get("/admin/orders")
        assert "Contacted" not in orders.text
        assert "Negotiating" not in orders.text
        assert "Request From" in orders.text and "Request To" in orders.text
        assert "Requested" in orders.text
        dashboard = client.get("/admin")
        assert "Requested" in dashboard.text

        detail = client.get("/admin/orders/1")
        assert detail.status_code == 200
        assert "Day" in detail.text
        assert detail.text.index("Selected Menu") < detail.text.index("Quote & Confirmation")
        assert detail.text.index("Quote & Confirmation") < detail.text.index("Kitchen Sheet & Sharing")
        assert detail.text.index("Kitchen Sheet & Sharing") < detail.text.index("Order Finance")
        assert "Web Order Charge" not in detail.text
        assert 'Set the status to Quoted first' in detail.text
        assert 'Change the order status to Confirmed' in detail.text
        csrf = csrf_from(detail.text)
        upd = client.post("/admin/orders/1", data={
            "csrf_token": csrf,
            "status": "quoted",
            "adult_charge": "20.00",
            "kid_charge": "10.00",
            "delivery_service_charge": "55.00",
            "admin_notes": "Kitchen note",
            "customer_message": "Quote ready.",
        }, follow_redirects=False)
        assert upd.status_code == 303

        customer = client.get("/orders/" + b1["token"])
        assert customer.status_code == 200
        assert "Delivery &amp; Service Charge" in customer.text
        assert "495.00" in customer.text
        assert "Want To Negotiate" not in customer.text
        assert "Invoice Available" not in customer.text
        assert "Cancel Order" not in customer.text
        assert "Wednesday" not in customer.text or "Day" in customer.text  # day is derived, date-dependent
        assert "+ Add More Dishes" not in customer.text
        assert "Please contact the Spice India Catering team" in customer.text
        assert "CONFIRM ORDER" in customer.text
        locked_edit = client.get("/orders/" + b1["token"] + "/add-dishes", follow_redirects=False)
        assert locked_edit.status_code == 303

        transaction = client.get("/admin/transactions/1")
        assert transaction.status_code == 200
        assert "Delivery & Service Charge" in transaction.text
        assert "Web Order" not in transaction.text

        transactions = client.get("/admin/transactions")
        assert 'data-transaction-filter-drawer' in transactions.text
        assert "Payment Status" in transactions.text
        assert "Expenses" in transactions.text
        assert "Sort By" in transactions.text

        # Kitchen actions are locked while Quoted.
        detail = client.get("/admin/orders/1")
        csrf = csrf_from(detail.text)
        kitchen_locked = client.post("/admin/orders/1/kitchen-share", data={
            "csrf_token": csrf,
            "kitchen_comments": "NO ONION IN 2 PORTIONS",
        }, follow_redirects=False)
        assert kitchen_locked.status_code == 303
        assert kitchen_locked.headers["location"].startswith("/admin/orders/1?toast=")
        assert client.get(f"/kitchen/{b1['token']}.pdf").status_code == 409

        # Once Confirmed, kitchen PDF/WhatsApp unlocks.
        detail = client.get("/admin/orders/1")
        confirm_admin = client.post("/admin/orders/1", data={
            "csrf_token": csrf_from(detail.text),
            "status": "confirmed",
            "adult_charge": "20.00",
            "kid_charge": "10.00",
            "delivery_service_charge": "55.00",
            "admin_notes": "Kitchen note",
            "customer_message": "Confirmed.",
        }, follow_redirects=False)
        assert confirm_admin.status_code == 303
        detail = client.get("/admin/orders/1")
        kitchen = client.post("/admin/orders/1/kitchen-share", data={
            "csrf_token": csrf_from(detail.text),
            "kitchen_comments": "NO ONION IN 2 PORTIONS",
        }, follow_redirects=False)
        assert kitchen.status_code == 303
        assert kitchen.headers["location"].startswith("https://wa.me/")
        decoded = unquote(kitchen.headers["location"])
        assert "/kitchen/" in decoded and ".pdf" in decoded
        kitchen_pdf = client.get(f"/kitchen/{b1['token']}.pdf")
        assert kitchen_pdf.status_code == 200
        assert kitchen_pdf.content.startswith(b"%PDF")

        # Menu administration is Excel-only.
        menus = client.get("/admin/menus", follow_redirects=False)
        assert menus.status_code == 303
        assert menus.headers["location"] == "/admin/menu-import"
        menu_import = client.get("/admin/menu-import")
        assert menu_import.status_code == 200
        assert "One Master Menu" in menu_import.text
        assert "Download Current Excel" in menu_import.text
        assert "Menu Images" in menu_import.text
        assert "Database backup active" in menu_import.text or "Cloudflare R2 backup connected" in menu_import.text
        assert "Master Category Setup" in menu_import.text
        master_download = client.get("/admin/menu-import/download")
        assert master_download.status_code == 200
        assert master_download.content[:2] == b"PK"

        settings = client.get("/admin/settings")
        assert "Web Order Charge" not in settings.text
        assert "Kitchen WhatsApp Number" in settings.text
        assert "Kitchen WhatsApp Group Link" in settings.text
        assert "Reset Test Orders & Restart Numbering" in settings.text
        settings_csrf = csrf_from(settings.text)
        reset = client.post("/admin/settings/reset-orders", data={
            "csrf_token": settings_csrf,
            "confirmation": "RESET ORDERS",
        }, follow_redirects=False)
        assert reset.status_code == 303
        after_reset_settings = client.get("/admin/settings")
        assert "Reset Test Orders & Restart Numbering" not in after_reset_settings.text
        reset_again = client.post("/admin/settings/reset-orders", data={
            "csrf_token": csrf_from(after_reset_settings.text),
            "confirmation": "RESET ORDERS",
        }, follow_redirects=False)
        assert reset_again.status_code == 303

        from sqlalchemy import func, select
        from app.models import QuoteRequest
        with SessionLocal() as db:
            assert (db.scalar(select(func.count(QuoteRequest.id))) or 0) == 0

        q3 = client.post("/api/quotes", json=quote_payload(event_name="After Reset", days=15))
        assert q3.status_code == 200, q3.text
        assert q3.json()["order_number"] == "CAT0001"


def teardown_module():
    from app.db import engine
    engine.dispose()
    if TEST_DB.exists():
        TEST_DB.unlink()


def test_excel_menu_import_and_combination_rules():
    from io import BytesIO
    from openpyxl import Workbook
    from sqlalchemy import select
    from app.db import SessionLocal
    from app.main import apply_menu_workbook
    from app.menu_excel import parse_menu_workbook
    from app.models import MenuCombination, MenuItem

    wb = Workbook()
    ws = wb.active
    ws.title = "Menu Import"
    ws.append(["Complete Menu"])
    ws.append([])
    ws.append(["Item ID", "Main Category", "South Indian", "North Indian", "Sub Category", "Menu Section", "Menu Item Name", "Item Image", "Display Price (€)", "Display Order", "Active", "Notes", "Food Description"])
    ws.append([501, "Veg Cuisine", "Yes", "Yes", "Bread's/Rice", "Indian Breads", "Test Appam", "https://example.com/appam.jpg", 4.50, 1, "Yes", "", "Soft South Indian bread served fresh."])
    ws.append([502, "Non Veg Cuisine", "Yes", "Yes", "Main", "Non-Vegetarian Main Course", "Test Curry", "", 9.95, 1, "Yes", "", "Rich house-style curry."])
    setup = wb.create_sheet("Category Setup")
    setup.append(["Website Category Structure"])
    setup.append([])
    setup.append(["Type", "Name", "Display Order", "Image / Icon"])
    setup.append(["Side Filter", "South Indian", 1, ""])
    setup.append(["Side Filter", "North Indian", 2, ""])
    setup.append(["Sub Category", "Bread's/Rice", 1, ""])
    setup.append(["Sub Category", "Main", 2, ""])
    sections = wb.create_sheet("Section Setup")
    sections.append(["Menu Section Headings"])
    sections.append([])
    sections.append(["Main Category", "Sub Category", "Menu Section", "Display Order", "Heading Colour", "Active"])
    sections.append(["Veg Cuisine", "Bread's/Rice", "Indian Breads", 1, "Green", "Yes"])
    sections.append(["Non Veg Cuisine", "Main", "Non-Vegetarian Main Course", 1, "Red", "Yes"])
    combos = wb.create_sheet("Combinations")
    combos.append(["Combination Rules"])
    combos.append([])
    combos.append(["Rule ID", "Trigger Item ID", "Trigger Item Name", "Recommended Item ID", "Recommended Item Name", "Priority", "Active", "Reciprocal", "Popup Title", "Notes"])
    combos.append([1, 501, "Test Appam", 502, "Test Curry", 1, "Yes", "Yes", "Popular combination", ""])
    bio = BytesIO(); wb.save(bio)

    parsed = parse_menu_workbook(bio.getvalue())
    assert len(parsed["items"]) == 2
    assert len(parsed["combinations"] or []) == 1
    with SessionLocal() as db:
        result = apply_menu_workbook(db, parsed)
        assert result["items"] == 2
        appam = db.scalar(select(MenuItem).where(MenuItem.import_item_id == 501).order_by(MenuItem.id))
        curry = db.scalar(select(MenuItem).where(MenuItem.import_item_id == 502).order_by(MenuItem.id))
        assert appam and curry
        assert float(appam.display_price) == 4.50
        assert appam.image_reference == "https://example.com/appam.jpg"
        assert appam.description == "Soft South Indian bread served fresh."
        assert appam.section_heading_color == "#22C55E"
        assert curry.section_heading_color == "#EF4444"
        assert db.scalar(select(MenuCombination).where(MenuCombination.trigger_import_item_id == 501))
        appam_id, curry_id = appam.id, curry.id

    with TestClient(app) as client:
        menu_page = client.get("/menu")
        assert menu_page.status_code == 200
        assert "floating-category-trigger" in menu_page.text
        assert "category-picker-dialog" in menu_page.text
        assert "https://example.com/appam.jpg" in menu_page.text
        assert '"price": 4.5' in menu_page.text or '"price":4.5' in menu_page.text
        js = client.get("/static/js/customer.js")
        assert "review-basket-image" in js.text
        response = client.get("/api/menu-combinations", params={"item_id": appam_id})
        assert response.status_code == 200
        assert response.json()["items"][0]["name"] == "Test Curry"
        response = client.get("/api/menu-combinations", params={"item_id": curry_id})
        assert response.status_code == 200
        assert response.json()["items"][0]["name"] == "Test Appam"
