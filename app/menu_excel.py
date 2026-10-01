from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook


YES_VALUES = {"yes", "y", "true", "1", "active"}
COLOUR_NAMES = {
    "green": "#22C55E",
    "red": "#EF4444",
    "teal": "#2DD4BF",
    "gold": "#F59E0B",
    "white": "#F8FAFC",
    "grey": "#94A3B8",
    "gray": "#94A3B8",
}


def _text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _yes(value: Any, default: bool = False) -> bool:
    text = _text(value).casefold()
    if not text:
        return default
    return text in YES_VALUES


def _int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return default


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        cleaned = str(value).replace("€", "").replace(",", "").strip()
        amount = Decimal(cleaned)
        if amount < 0:
            raise ValueError
        return amount.quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid price '{value}'. Use a positive number or leave it blank.")


def _colour(value: Any, default: str = "#94A3B8") -> str:
    raw = _text(value)
    if not raw:
        return default
    named = COLOUR_NAMES.get(raw.casefold())
    if named:
        return named
    if re.fullmatch(r"#[0-9A-Fa-f]{6}", raw):
        return raw.upper()
    if re.fullmatch(r"#[0-9A-Fa-f]{3}", raw):
        return "#" + "".join(ch * 2 for ch in raw[1:]).upper()
    return default




def _cell_colour(cell, default: str) -> str:
    typed = _text(cell.value)
    if typed:
        return _colour(typed, default)
    try:
        fill = cell.fill
        fg = fill.fgColor
        rgb = str(getattr(fg, "rgb", "") or "")
        if fill.fill_type and len(rgb) >= 6:
            return _colour("#" + rgb[-6:], default)
    except Exception:
        pass
    return default

def _find_header(ws, required: str, max_rows: int = 12) -> tuple[int, dict[str, int]]:
    for row in range(1, min(ws.max_row, max_rows) + 1):
        headers: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = _text(ws.cell(row, col).value)
            if value:
                headers[value] = col
        if required in headers:
            return row, headers
    raise ValueError(f"Could not find the '{required}' header in sheet '{ws.title}'.")


def _embedded_images_by_row(ws, image_col: int, header_row: int) -> dict[int, tuple[bytes, str]]:
    """Return images anchored in the Item Image column, keyed by 1-based worksheet row."""
    result: dict[int, tuple[bytes, str]] = {}
    for image in getattr(ws, "_images", []):
        try:
            anchor = image.anchor._from
            row = int(anchor.row) + 1
            col = int(anchor.col) + 1
            if row <= header_row or col != image_col:
                continue
            raw = image._data()
            fmt = str(getattr(image, "format", "png") or "png").lower()
            content_type = "image/jpeg" if fmt in {"jpg", "jpeg"} else f"image/{fmt}"
            result[row] = (raw, content_type)
        except Exception:
            continue
    return result


def parse_menu_workbook(data: bytes) -> dict[str, Any]:
    """Parse the Excel-only catering menu workbook without mutating the database."""
    try:
        wb = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc

    if "Menu Import" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Menu Import' sheet.")

    ws = wb["Menu Import"]
    header_row, headers = _find_header(ws, "Item ID")
    required = ["Item ID", "Main Category", "Sub Category", "Menu Section", "Menu Item Name", "Display Order", "Active"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("Menu Import is missing columns: " + ", ".join(missing))

    # Region columns are intentionally dynamic: every column between Main Category and Sub Category.
    main_col = headers["Main Category"]
    sub_col = headers["Sub Category"]
    region_columns: list[tuple[str, int]] = []
    for col in range(main_col + 1, sub_col):
        name = _text(ws.cell(header_row, col).value)
        if name:
            region_columns.append((name, col))

    image_col = headers.get("Item Image")
    embedded_images = _embedded_images_by_row(ws, image_col, header_row) if image_col else {}
    price_col = headers.get("Display Price (€)") or headers.get("Price (€)") or headers.get("Price")

    items: list[dict[str, Any]] = []
    seen_ids: set[int] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        raw_id = ws.cell(row, headers["Item ID"]).value
        name = _text(ws.cell(row, headers["Menu Item Name"]).value)
        if raw_id in (None, "") and not name:
            continue
        item_id = _int(raw_id, -1)
        if item_id <= 0:
            raise ValueError(f"Menu Import row {row}: Item ID must be a positive whole number.")
        if item_id in seen_ids:
            raise ValueError(f"Menu Import row {row}: duplicate Item ID {item_id}.")
        seen_ids.add(item_id)
        if not name:
            raise ValueError(f"Menu Import row {row}: Menu Item Name is required.")

        main_category = _text(ws.cell(row, headers["Main Category"]).value)
        main_key = main_category.casefold()
        if "non" in main_key and "veg" in main_key:
            dietary = "nonveg"
        elif main_key == "both" or "both" in main_key:
            dietary = "both"
        else:
            dietary = "veg"

        regions = [region_name for region_name, col in region_columns if _yes(ws.cell(row, col).value, False)]
        if not regions:
            regions = [region_name for region_name, _ in region_columns]

        try:
            display_price = _decimal(ws.cell(row, price_col).value) if price_col else None
        except ValueError as exc:
            raise ValueError(f"Menu Import row {row}: {exc}") from exc

        embedded = embedded_images.get(row)
        items.append({
            "item_id": item_id,
            "main_category": main_category,
            "dietary": dietary,
            "regions": regions,
            "category": _text(ws.cell(row, headers["Sub Category"]).value) or "Other",
            "section": _text(ws.cell(row, headers["Menu Section"]).value) or "Main Selection",
            "name": name,
            "description": _text(ws.cell(row, headers.get("Food Description", 0)).value) if headers.get("Food Description") else (_text(ws.cell(row, headers.get("Notes", 0)).value) if headers.get("Notes") else ""),
            "image_reference": _text(ws.cell(row, image_col).value) if image_col else "",
            "image_bytes": embedded[0] if embedded else None,
            "image_content_type": embedded[1] if embedded else None,
            "display_price": display_price,
            "sort_order": _int(ws.cell(row, headers["Display Order"]).value, 0),
            "active": _yes(ws.cell(row, headers["Active"]).value, True),
            "notes": _text(ws.cell(row, headers.get("Notes", 0)).value) if headers.get("Notes") else "",
        })

    side_filters: list[dict[str, Any]] = []
    category_setup: dict[str, dict[str, Any]] = {}
    if "Category Setup" in wb.sheetnames:
        setup = wb["Category Setup"]
        try:
            setup_header, sh = _find_header(setup, "Type")
            for row in range(setup_header + 1, setup.max_row + 1):
                type_name = _text(setup.cell(row, sh["Type"]).value).casefold()
                name = _text(setup.cell(row, sh.get("Name", 0)).value)
                if not name:
                    continue
                if type_name == "side filter":
                    side_filters.append({
                        "name": name,
                        "sort_order": _int(setup.cell(row, sh.get("Display Order", 0)).value, 0),
                    })
                elif type_name == "sub category":
                    category_setup[name.casefold()] = {
                        "name": name,
                        "sort_order": _int(setup.cell(row, sh.get("Display Order", 0)).value, 0),
                    }
        except ValueError:
            pass
    if not side_filters:
        side_filters = [{"name": name, "sort_order": idx + 1} for idx, (name, _) in enumerate(region_columns)]
    side_filters = [f for f in side_filters if f["name"]]

    section_setup: dict[tuple[str, str, str], dict[str, Any]] = {}
    if "Section Setup" in wb.sheetnames:
        sw = wb["Section Setup"]
        try:
            section_header, sh = _find_header(sw, "Menu Section")
            required_section = ["Main Category", "Sub Category", "Menu Section", "Display Order", "Heading Colour", "Active"]
            missing_section = [name for name in required_section if name not in sh]
            if missing_section:
                raise ValueError("Section Setup is missing columns: " + ", ".join(missing_section))
            for row in range(section_header + 1, sw.max_row + 1):
                main = _text(sw.cell(row, sh["Main Category"]).value)
                category = _text(sw.cell(row, sh["Sub Category"]).value)
                section = _text(sw.cell(row, sh["Menu Section"]).value)
                if not section:
                    continue
                default_colour = "#EF4444" if "non" in main.casefold() and "veg" in main.casefold() else "#22C55E" if "veg" in main.casefold() else "#2DD4BF"
                section_setup[(main.casefold(), category.casefold(), section.casefold())] = {
                    "sort_order": _int(sw.cell(row, sh["Display Order"]).value, 0),
                    "heading_color": _cell_colour(sw.cell(row, sh["Heading Colour"]), default_colour),
                    "active": _yes(sw.cell(row, sh["Active"]).value, True),
                }
        except ValueError as exc:
            raise exc

    combinations: list[dict[str, Any]] | None = None
    if "Combinations" in wb.sheetnames:
        combinations = []
        cw = wb["Combinations"]
        combo_header, ch = _find_header(cw, "Trigger Item ID")
        required_combo = ["Trigger Item ID", "Recommended Item ID", "Priority", "Active", "Reciprocal"]
        missing_combo = [name for name in required_combo if name not in ch]
        if missing_combo:
            raise ValueError("Combinations is missing columns: " + ", ".join(missing_combo))
        for row in range(combo_header + 1, cw.max_row + 1):
            trigger_raw = cw.cell(row, ch["Trigger Item ID"]).value
            rec_raw = cw.cell(row, ch["Recommended Item ID"]).value
            if trigger_raw in (None, "") and rec_raw in (None, ""):
                continue
            trigger = _int(trigger_raw, -1)
            recommended = _int(rec_raw, -1)
            if trigger <= 0 or recommended <= 0:
                raise ValueError(f"Combinations row {row}: Trigger and Recommended Item IDs must be positive numbers.")
            if trigger == recommended:
                raise ValueError(f"Combinations row {row}: an item cannot recommend itself.")
            if trigger not in seen_ids:
                raise ValueError(f"Combinations row {row}: Trigger Item ID {trigger} is not in Menu Import.")
            if recommended not in seen_ids:
                raise ValueError(f"Combinations row {row}: Recommended Item ID {recommended} is not in Menu Import.")
            combinations.append({
                "trigger_item_id": trigger,
                "recommended_item_id": recommended,
                "priority": max(1, _int(cw.cell(row, ch["Priority"]).value, 1)),
                "active": _yes(cw.cell(row, ch["Active"]).value, True),
                "reciprocal": _yes(cw.cell(row, ch["Reciprocal"]).value, False),
                "popup_title": _text(cw.cell(row, ch.get("Popup Title", 0)).value) if ch.get("Popup Title") else "",
                "notes": _text(cw.cell(row, ch.get("Notes", 0)).value) if ch.get("Notes") else "",
            })

    return {
        "items": items,
        "side_filters": side_filters,
        "category_setup": category_setup,
        "section_setup": section_setup,
        "combinations": combinations,
        "region_columns": [name for name, _ in region_columns],
    }


def update_workbook_image_references(data: bytes, references: dict[int, str]) -> bytes:
    """Return a workbook copy with Item Image cells updated by stable Item ID."""
    if not references:
        return data
    try:
        wb = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError("The stored master workbook is not readable.") from exc
    if "Menu Import" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Menu Import' sheet.")
    ws = wb["Menu Import"]
    header_row, headers = _find_header(ws, "Item ID")
    image_col = headers.get("Item Image")
    if not image_col:
        raise ValueError("Menu Import must contain an 'Item Image' column.")
    for row in range(header_row + 1, ws.max_row + 1):
        item_id = _int(ws.cell(row, headers["Item ID"]).value, -1)
        if item_id in references:
            ws.cell(row, image_col).value = references[item_id]
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def update_workbook_category_setup(
    data: bytes,
    *,
    old_name: str = "",
    new_name: str,
    display_order: int,
) -> bytes:
    """Rename/add a master Sub Category and keep Menu Import/Section Setup in sync.

    The stored Excel remains the source of truth, so an Admin category edit must also
    update the workbook or the next Excel import would undo the website change.
    """
    try:
        wb = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError("The stored master workbook is not readable.") from exc

    clean_new = _text(new_name)
    clean_old = _text(old_name)
    if not clean_new:
        raise ValueError("Category name is required.")
    if display_order < 1:
        raise ValueError("Display order must be 1 or greater.")

    if "Category Setup" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Category Setup' sheet.")
    setup = wb["Category Setup"]
    setup_header, sh = _find_header(setup, "Type")
    if "Name" not in sh or "Display Order" not in sh:
        raise ValueError("Category Setup must contain Name and Display Order columns.")

    target_row = None
    lookup = clean_old.casefold() if clean_old else clean_new.casefold()
    for row in range(setup_header + 1, setup.max_row + 1):
        type_name = _text(setup.cell(row, sh["Type"]).value).casefold()
        name = _text(setup.cell(row, sh["Name"]).value)
        if type_name == "sub category" and name.casefold() == lookup:
            target_row = row
            break

    if target_row is None:
        # Append after the last configured row. Blank/formatted rows are fine; using
        # max_row + 1 avoids overwriting existing workbook notes.
        target_row = setup.max_row + 1
        setup.cell(target_row, sh["Type"]).value = "Sub Category"
    setup.cell(target_row, sh["Name"]).value = clean_new
    setup.cell(target_row, sh["Display Order"]).value = int(display_order)

    if clean_old and clean_old.casefold() != clean_new.casefold():
        if "Menu Import" in wb.sheetnames:
            ws = wb["Menu Import"]
            header_row, headers = _find_header(ws, "Item ID")
            sub_col = headers.get("Sub Category")
            if sub_col:
                for row in range(header_row + 1, ws.max_row + 1):
                    if _text(ws.cell(row, sub_col).value).casefold() == clean_old.casefold():
                        ws.cell(row, sub_col).value = clean_new

        if "Section Setup" in wb.sheetnames:
            sw = wb["Section Setup"]
            section_header, headers = _find_header(sw, "Menu Section")
            sub_col = headers.get("Sub Category")
            if sub_col:
                for row in range(section_header + 1, sw.max_row + 1):
                    if _text(sw.cell(row, sub_col).value).casefold() == clean_old.casefold():
                        sw.cell(row, sub_col).value = clean_new

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def parse_internal_price_workbook(data: bytes) -> list[dict[str, Any]]:
    """Parse the admin-only internal dish price workbook.

    Expected columns: Item Identifier, Item Name, Price.
    Price is treated as an internal cost per guest for catering costing.
    """
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc
    ws = wb["Internal Item Prices"] if "Internal Item Prices" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row, headers = _find_header(ws, "Item Identifier", max_rows=20)
    required = ["Item Identifier", "Item Name", "Price"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("Internal price workbook is missing columns: " + ", ".join(missing))
    rows: list[dict[str, Any]] = []
    seen: set[int] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        raw_id = ws.cell(row, headers["Item Identifier"]).value
        name = _text(ws.cell(row, headers["Item Name"]).value)
        raw_price = ws.cell(row, headers["Price"]).value
        if raw_id in (None, "") and not name and raw_price in (None, ""):
            continue
        item_id = _int(raw_id, -1)
        if item_id <= 0:
            raise ValueError(f"Internal price row {row}: Item Identifier must be a positive whole number.")
        if item_id in seen:
            raise ValueError(f"Internal price row {row}: duplicate Item Identifier {item_id}.")
        seen.add(item_id)
        if not name:
            raise ValueError(f"Internal price row {row}: Item Name is required.")
        try:
            price = _decimal(raw_price)
        except ValueError as exc:
            raise ValueError(f"Internal price row {row}: {exc}") from exc
        rows.append({"item_id": item_id, "name": name, "price": price})
    if not rows:
        raise ValueError("The internal price workbook does not contain any menu items.")
    return rows


def build_internal_price_workbook(items: list[dict[str, Any]]) -> bytes:
    """Build a clean three-column admin-only internal price workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Internal Item Prices"
    ws.append(["Item Identifier", "Item Name", "Price"])
    for item in items:
        ws.append([item.get("item_id"), item.get("name", ""), item.get("price")])
    fill = PatternFill("solid", fgColor="0B3B36")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    for cell in ws["C"][1:]:
        cell.number_format = '€0.00'
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
